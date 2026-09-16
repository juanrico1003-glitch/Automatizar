"""
Script de Automatización Integral para Cuentas por Pagar (CXP)
Replica al 100% y supera la funcionalidad de la Macro VBA:
- Extrae Proveedor, Cuenta Contable, Documento, Fechas y Valores desde el reporte crudo.
- Configura los 36 encabezados completos en la Fila 5.
- Calcula FECHA HELISA (=F.VENCE - TIEMPO CREDITO).
- Calcula TIEMPO CREDITO (=F.VENCE - FECHA).
- Distribuye automáticamente los valores en las 24 Quincenas del año (ENE Q1 hasta DIC Q2).
- Calcula los subtotales de cada quincena en la Fila 4 con fórmulas SUM.
- Aplica formato de moneda, formato de fechas, anchos de columna, autofiltro y paneles congelados.
- Compatible tanto con archivos .xlsx como .xlsm.
"""

import sys
import os
import glob
import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Mapeo de abreviaturas de meses en español
MESES_MAP = {
    'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6,
    'JUL': 7, 'AGO': 8, 'SEP': 9, 'OCT': 10, 'NOV': 11, 'DIC': 12
}


def parse_fecha(dato):
    """Convierte texto de fecha (ej. '24/AGO/2026' o '24/08/2026') en objeto date de Python."""
    if dato is None:
        return None
    if isinstance(dato, datetime.datetime):
        return dato.date()
    if isinstance(dato, datetime.date):
        return dato
    texto = str(dato).strip().upper()
    if not texto:
        return None
    partes = texto.split('/')
    if len(partes) == 3:
        try:
            d = int(partes[0])
            m_str = partes[1]
            y = int(partes[2])
            if m_str in MESES_MAP:
                return datetime.date(y, MESES_MAP[m_str], d)
            elif m_str.isdigit():
                return datetime.date(y, int(m_str), d)
        except Exception:
            return None
    return None


def procesar_archivo_cxp(filepath):
    print(f"\n=======================================================")
    print(f" Procesando: {os.path.basename(filepath)}")
    print(f"=======================================================")

    is_xlsm = filepath.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(filepath, keep_vba=is_xlsm)
    
    # 1. Buscar la hoja de origen
    raw_sheet_name = None
    for name in wb.sheetnames:
        if "DETALLADO POR EDADES" in name.upper():
            raw_sheet_name = name
            break
            
    if not raw_sheet_name:
        print(f"ERROR: No se encontró la hoja con el reporte crudo ('DETALLADO POR EDADES') en {filepath}")
        return False

    raw_sheet = wb[raw_sheet_name]
    
    # 2. Obtener o crear Hoja1
    if "Hoja1" in wb.sheetnames:
        out_sheet = wb["Hoja1"]
    else:
        out_sheet = wb.create_sheet(title="Hoja1")

    # Limpiar completamente cualquier dato anterior en Hoja1
    if out_sheet.max_row >= 1:
        out_sheet.delete_rows(1, out_sheet.max_row + 10)

    # 3. Definir los 36 encabezados
    encabezados = [
        'REVISAR', 'vto oc', 'CTA', 'PROVEEDOR', 'No. FC', 'FECHA HELISA', 'FECHA', 'F.VENCE',
        'valor', 'TIEMPO CREDITO', 'ENE Q1', 'ENE Q2', 'FEB Q1', 'FEB Q2', 'MAR Q1', 'MAR Q2',
        'ABR Q1', 'ABR Q2', 'MAY Q1', 'MAY Q2', 'JUN Q1', 'JUN Q2', 'JUL Q1', 'JUL Q2',
        'AGO Q1', 'AGO Q2', 'SEP Q1', 'SEP Q2', 'OCT Q1', 'OCT Q2', 'NOV Q1', 'NOV Q2',
        'DIC Q1', 'DIC Q2', 'pago', 'VALOR'
    ]

    header_font = Font(name='Aptos Narrow', size=11, bold=True, color='000000')
    for col_idx, h_text in enumerate(encabezados, start=1):
        c = out_sheet.cell(5, col_idx, h_text)
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    # 4. Extraer registros del reporte crudo
    curr_supp = None
    curr_cta = None
    extracted_rows = []
    suppliers_set = set()
    total_valor = 0.0

    for r in range(1, raw_sheet.max_row + 1):
        val_a = raw_sheet.cell(r, 1).value
        val_f = raw_sheet.cell(r, 6).value   # FECHA
        val_j = raw_sheet.cell(r, 10).value  # Vr. Total
        val_o = raw_sheet.cell(r, 15).value  # F.VENCE

        # Identificar encabezado de Proveedor o Cuenta en columna A
        if val_a and isinstance(val_a, str):
            val_a_str = val_a.strip()
            if any(h in val_a_str for h in ["SODECA LATAM", "Estado detallado", "CUENTA "]):
                continue
            
            if "(Dir:" in val_a_str:
                curr_supp = val_a_str
                curr_cta = None
                suppliers_set.add(curr_supp)
                continue
                
            tokens = val_a_str.split()
            if tokens and tokens[0].isdigit() and len(tokens[0]) >= 4:
                curr_cta = val_a_str
                continue

        # Identificar registros de transacciones (con monto numérico en Col J)
        if val_j is not None and isinstance(val_j, (int, float)):
            doc_str = str(val_a).strip() if val_a is not None else None
            fecha_doc = parse_fecha(val_f)
            fecha_vence = parse_fecha(val_o)
            valor = float(val_j)
            
            extracted_rows.append((curr_cta, curr_supp, doc_str, fecha_doc, fecha_vence, valor))
            total_valor += valor

    if not extracted_rows:
        print("No se encontraron registros para organizar.")
        return False

    # 5. Escribir datos, fórmulas y formatos en Hoja1 a partir de la fila 6
    fila_salida = 6
    data_font = Font(name='Aptos Narrow', size=11, bold=False, color='000000')
    currency_fmt = '"$" #,##0.00;-"$" #,##0.00'
    date_fmt = 'dd/mmm/yyyy'

    for idx, (cta, supp, doc, fecha_doc, fecha_vence, valor) in enumerate(extracted_rows):
        r_out = fila_salida + idx
        
        # Col A, B: REVISAR, vto oc (vacíos)
        out_sheet.cell(r_out, 1).value = None
        out_sheet.cell(r_out, 2).value = None
        
        # Col C: CTA
        out_sheet.cell(r_out, 3).value = cta
        
        # Col D: PROVEEDOR
        out_sheet.cell(r_out, 4).value = supp
        
        # Col E: No. FC
        if doc:
            c_e = out_sheet.cell(r_out, 5)
            c_e.value = doc
            c_e.alignment = Alignment(horizontal='center')
        else:
            out_sheet.cell(r_out, 5).value = None

        # Col F: FECHA HELISA (=IF(OR(H6="",J6=""),"",H6-J6))
        c_f = out_sheet.cell(r_out, 6)
        c_f.value = f'=IF(OR(H{r_out}="",J{r_out}=""),"",H{r_out}-J{r_out})'
        c_f.number_format = date_fmt
        c_f.alignment = Alignment(horizontal='center')

        # Col G: FECHA
        if fecha_doc:
            c_g = out_sheet.cell(r_out, 7)
            c_g.value = fecha_doc
            c_g.number_format = date_fmt
            c_g.alignment = Alignment(horizontal='center')
        else:
            out_sheet.cell(r_out, 7).value = None

        # Col H: F.VENCE
        if fecha_vence:
            c_h = out_sheet.cell(r_out, 8)
            c_h.value = fecha_vence
            c_h.number_format = date_fmt
            c_h.alignment = Alignment(horizontal='center')
        else:
            out_sheet.cell(r_out, 8).value = None

        # Col I: valor
        c_i = out_sheet.cell(r_out, 9)
        c_i.value = valor
        c_i.number_format = currency_fmt

        # Col J: TIEMPO CREDITO (=IF(OR(G6="",H6=""),"",H6-G6))
        c_j = out_sheet.cell(r_out, 10)
        c_j.value = f'=IF(OR(G{r_out}="",H{r_out}=""),"",H{r_out}-G{r_out})'
        c_j.number_format = '0'
        c_j.alignment = Alignment(horizontal='center')

        # Col K a AH: 24 Quincenas del año (Mes 1 a 12, Q1 y Q2)
        for mes in range(1, 13):
            col_q1 = 11 + (mes - 1) * 2
            col_q2 = col_q1 + 1

            # Q1 (días 1 al 15)
            cq1 = out_sheet.cell(r_out, col_q1)
            cq1.value = f'=IF(AND($G{r_out}<>"",MONTH($G{r_out})={mes},DAY($G{r_out})<=15),$I{r_out},"")'
            cq1.number_format = currency_fmt

            # Q2 (días 16 al fin de mes)
            cq2 = out_sheet.cell(r_out, col_q2)
            cq2.value = f'=IF(AND($G{r_out}<>"",MONTH($G{r_out})={mes},DAY($G{r_out})>=16),$I{r_out},"")'
            cq2.number_format = currency_fmt

        # Col AI, AJ: pago, VALOR (vacíos)
        out_sheet.cell(r_out, 35).value = None
        out_sheet.cell(r_out, 36).value = None

        for c_idx in range(1, 37):
            out_sheet.cell(r_out, c_idx).font = data_font

    last_row = fila_salida + len(extracted_rows) - 1

    # 6. Subtotales de cada quincena en la Fila 4 (K4:AH4)
    sum_font = Font(name='Aptos Narrow', size=11, bold=True, color='000000')
    for col_idx in range(11, 35):
        col_let = get_column_letter(col_idx)
        c_sum = out_sheet.cell(4, col_idx)
        c_sum.value = f'=SUM({col_let}6:{col_let}{last_row})'
        c_sum.font = sum_font
        c_sum.number_format = currency_fmt

    # 7. Configuración estética: Anchos de columna, Autofiltro y Paneles
    col_widths = {
        'A': 12.8, 'B': 13.0, 'C': 24.8, 'D': 55.8, 'E': 18.8,
        'F': 14.8, 'G': 13.0, 'H': 13.0, 'I': 16.8, 'J': 17.8,
        'AI': 14.8, 'AJ': 13.0
    }
    for col_let, w in col_widths.items():
        out_sheet.column_dimensions[col_let].width = w
    for col_idx in range(11, 35):
        out_sheet.column_dimensions[get_column_letter(col_idx)].width = 13.5

    out_sheet.auto_filter.ref = f'A5:AJ{last_row}'
    out_sheet.freeze_panes = 'A6'

    wb.save(filepath)
    print(f" -> Proveedores procesados: {len(suppliers_set)}")
    print(f" -> Registros organizados:  {len(extracted_rows)}")
    print(f" -> Rango de datos:         Fila 6 a Fila {last_row}")
    print(f" -> Columnas configuradas:  36 columnas (incluyendo 24 quincenas y subtotales en fila 4)")
    print(f" -> Suma total de valores:  ${total_valor:,.2f}")
    print(f" -> Archivo guardado exitosamente con cálculo completo.")
    return True


def main():
    if len(sys.argv) > 1:
        target_files = [sys.argv[1]]
    else:
        # Busca archivos Excel en la carpeta
        target_files = glob.glob("*.xlsx") + glob.glob("*.xlsm")
        # Excluir temporales de Excel o copias de referencia/respaldo
        target_files = [
            f for f in target_files 
            if not os.path.basename(f).startswith("~$") 
            and not f.endswith("_original.xlsx")
            and not "copia a seguir" in f.lower()
        ]

    if not target_files:
        print("No se encontraron archivos Excel para procesar en esta carpeta.")
        return

    for f in target_files:
        procesar_archivo_cxp(f)

    print("\n¡Proceso terminado exitosamente para todos los archivos!")


if __name__ == "__main__":
    main()
